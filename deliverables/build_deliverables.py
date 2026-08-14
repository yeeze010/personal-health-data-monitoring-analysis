from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


BASE = Path(__file__).resolve().parent
DIAGRAM_DIR = BASE / "diagrams" / "personal-health-platform"
DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)


project = {
    "name": "个人健康生活网络数据监测与分析系统",
    "stack": "Vue 3 + TypeScript、FastAPI、PostgreSQL、Redis、MinIO、Docker + Nginx",
    "summary": "面向个人、家庭健康管理员、医生/健康顾问和平台运营人员的健康数据采集、监测、分析、预警与报告系统。",
}


modules = [
    ("认证与权限", "登录、注册、角色权限、组织/家庭成员授权、操作审计"),
    ("用户与健康档案", "个人资料、基础指标、既往史、用药史、过敏史、紧急联系人"),
    ("设备与数据接入", "手动录入、CSV 导入、可穿戴设备 API 预留、数据校验"),
    ("健康指标管理", "血压、血糖、心率、体重、睡眠、运动、体温、血氧等指标"),
    ("趋势分析与看板", "多维趋势图、异常波动、目标达成、周期统计、家庭成员对比"),
    ("风险评估与预警", "阈值规则、组合规则、风险等级、站内提醒、邮件/短信预留"),
    ("健康报告", "周报/月报、医生查看摘要、PDF/图片导出、历史报告归档"),
    ("建议与计划", "健康目标、饮食/运动/复测建议、执行打卡、计划调整"),
    ("消息通知", "提醒模板、通知记录、已读未读、失败重试"),
    ("文件与附件", "体检报告、化验单、处方图片、MinIO 文件管理"),
    ("系统管理后台", "用户管理、角色管理、指标字典、预警规则、系统参数"),
    ("审计与合规", "访问日志、数据变更日志、导出记录、隐私授权记录"),
]

pages = [
    ("前台", "登录/注册", "用户身份认证、找回密码、隐私协议确认"),
    ("前台", "健康总览", "今日指标、风险提醒、目标进度、最近报告入口"),
    ("前台", "指标录入", "按指标录入、批量录入、异常值即时提示"),
    ("前台", "趋势分析", "折线图、区间筛选、指标叠加、异常点解释"),
    ("前台", "健康报告", "报告列表、报告详情、导出下载"),
    ("前台", "健康计划", "目标设置、计划执行、打卡记录"),
    ("前台", "家庭成员", "成员授权、成员切换、查看范围设置"),
    ("后台", "运营看板", "用户增长、数据接入量、预警量、系统健康状态"),
    ("后台", "用户管理", "用户查询、冻结/解冻、角色分配、档案查看"),
    ("后台", "指标字典", "指标单位、正常范围、录入约束、展示顺序"),
    ("后台", "预警规则", "阈值规则、组合规则、启停、灰度发布"),
    ("后台", "文件管理", "附件检索、下载、删除、存储占用统计"),
    ("后台", "审计日志", "操作记录、访问记录、导出记录、风险行为筛查"),
    ("后台", "系统设置", "通知渠道、存储配置、隐私策略、备份策略"),
]

tables = [
    ("users", "用户主表", "id, email, phone, password_hash, status, last_login_at, created_at"),
    ("roles", "角色表", "id, code, name, description"),
    ("user_roles", "用户角色关系", "id, user_id, role_id"),
    ("health_profiles", "健康档案", "id, user_id, gender, birthday, height_cm, blood_type, medical_history"),
    ("family_members", "家庭成员授权", "id, owner_user_id, member_user_id, relation, scope, status"),
    ("metric_definitions", "指标字典", "id, code, name, unit, data_type, normal_min, normal_max"),
    ("health_records", "健康数据记录", "id, user_id, metric_code, value_numeric, value_text, measured_at, source"),
    ("device_sources", "设备数据源", "id, user_id, provider, external_user_id, token_ref, status"),
    ("risk_rules", "预警规则", "id, metric_code, operator, threshold_value, risk_level, enabled"),
    ("risk_events", "风险事件", "id, user_id, rule_id, record_id, level, status, handled_at"),
    ("health_reports", "健康报告", "id, user_id, report_type, period_start, period_end, summary, file_id"),
    ("health_plans", "健康计划", "id, user_id, goal_type, target_value, start_date, end_date, status"),
    ("plan_checkins", "计划打卡", "id, plan_id, checkin_date, value, note"),
    ("files", "文件表", "id, owner_user_id, bucket, object_key, file_name, mime_type, size_bytes"),
    ("notifications", "通知记录", "id, user_id, channel, template_code, content, status, sent_at"),
    ("audit_logs", "审计日志", "id, actor_user_id, action, resource_type, resource_id, ip, created_at"),
]

apis = [
    ("POST", "/api/v1/auth/login", "登录并返回访问令牌"),
    ("GET", "/api/v1/me/profile", "读取个人健康档案"),
    ("PUT", "/api/v1/me/profile", "更新个人健康档案"),
    ("GET", "/api/v1/metrics", "读取指标字典"),
    ("POST", "/api/v1/records", "新增健康指标记录"),
    ("GET", "/api/v1/records", "分页查询健康指标记录"),
    ("POST", "/api/v1/records/import", "导入健康数据 CSV"),
    ("GET", "/api/v1/analytics/trends", "查询趋势分析数据"),
    ("GET", "/api/v1/analytics/summary", "查询健康总览统计"),
    ("GET", "/api/v1/risks", "查询风险事件"),
    ("PATCH", "/api/v1/risks/{id}/handle", "处理风险事件"),
    ("GET", "/api/v1/reports", "查询健康报告"),
    ("POST", "/api/v1/reports", "生成健康报告"),
    ("POST", "/api/v1/files/upload-url", "获取 MinIO 上传地址"),
    ("GET", "/api/v1/admin/users", "后台用户列表"),
    ("PUT", "/api/v1/admin/risk-rules/{id}", "后台维护预警规则"),
]

milestones = [
    ("M0 立项与需求冻结", "第 1 周", "需求规格、原型范围、验收口径确认"),
    ("M1 基础架构与权限", "第 2-3 周", "前后端脚手架、认证、权限、数据库迁移"),
    ("M2 健康数据核心闭环", "第 4-6 周", "档案、指标、录入、查询、导入、文件上传"),
    ("M3 分析预警与报告", "第 7-9 周", "趋势看板、风险规则、报告生成、通知记录"),
    ("M4 后台管理与联调", "第 10-11 周", "后台用户、字典、规则、审计、端到端联调"),
    ("M5 测试验收与上线", "第 12 周", "UAT、性能安全检查、部署交付、验收归档"),
]


def esc(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def write_svg(path: Path, body: str, viewbox: str = "0 0 1200 760"):
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="{viewbox}">
  <defs>
    <pattern id="grid" width="40" height="40" patternUnits="userSpaceOnUse">
      <path d="M 40 0 L 0 0 0 40" fill="none" stroke="#1e293b" stroke-width="0.5"/>
    </pattern>
    <marker id="arrow" markerWidth="10" markerHeight="7" refX="9" refY="3.5" orient="auto">
      <polygon points="0 0, 10 3.5, 0 7" fill="#94a3b8"/>
    </marker>
    <marker id="arrow-cyan" markerWidth="10" markerHeight="7" refX="9" refY="3.5" orient="auto">
      <polygon points="0 0, 10 3.5, 0 7" fill="#22d3ee"/>
    </marker>
    <style>
      @import url('https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@400;500;600;700&amp;display=swap');
      text {{ font-family: 'JetBrains Mono','Noto Sans SC','PingFang SC',sans-serif; }}
      .title {{ fill:#f8fafc; font-size:18px; font-weight:700; }}
      .sub {{ fill:#94a3b8; font-size:9px; }}
      .label {{ fill:#f8fafc; font-size:12px; font-weight:600; text-anchor:middle; }}
      .small {{ fill:#94a3b8; font-size:9px; text-anchor:middle; }}
      .tiny {{ fill:#cbd5e1; font-size:8px; }}
      .region {{ fill:none; stroke:#fbbf24; stroke-width:1; stroke-dasharray:8,4; }}
      .line {{ fill:none; stroke:#94a3b8; stroke-width:1.4; marker-end:url(#arrow); }}
      .line2 {{ fill:none; stroke:#22d3ee; stroke-width:1.6; marker-end:url(#arrow-cyan); }}
    </style>
  </defs>
  <rect width="100%" height="100%" fill="#0f172a"/>
  <rect width="100%" height="100%" fill="url(#grid)"/>
{body}
</svg>'''
    path.write_text(svg, encoding="utf-8")


def box(x, y, w, h, title, sub, fill, stroke):
    return f'''  <rect x="{x}" y="{y}" width="{w}" height="{h}" rx="6" fill="#0f172a"/>
  <rect x="{x}" y="{y}" width="{w}" height="{h}" rx="6" fill="{fill}" stroke="{stroke}" stroke-width="1.5"/>
  <text x="{x+w/2}" y="{y+25}" class="label">{esc(title)}</text>
  <text x="{x+w/2}" y="{y+43}" class="small">{esc(sub)}</text>\n'''


def cylinder(x, y, title, sub):
    return f'''  <g transform="translate({x},{y})">
    <rect x="0" y="12" width="145" height="56" fill="#0f172a"/>
    <ellipse cx="72.5" cy="12" rx="72.5" ry="12" fill="#0f172a"/>
    <rect x="0" y="12" width="145" height="56" fill="rgba(76,29,149,0.4)"/>
    <ellipse cx="72.5" cy="12" rx="72.5" ry="12" fill="rgba(76,29,149,0.4)" stroke="#a78bfa" stroke-width="1.5"/>
    <path d="M0,12 L0,68 M145,12 L145,68" stroke="#a78bfa" stroke-width="1.5"/>
    <ellipse cx="72.5" cy="68" rx="72.5" ry="12" fill="none" stroke="#a78bfa" stroke-width="1.5"/>
    <text x="72.5" y="38" class="label">{esc(title)}</text>
    <text x="72.5" y="54" class="small">{esc(sub)}</text>
  </g>\n'''


architecture = f'''  <text x="30" y="36" class="title">个人健康生活网络数据监测与分析系统｜系统架构图</text>
  <text x="30" y="56" class="sub">Docker + Nginx + FastAPI + PostgreSQL + Redis + MinIO</text>
  <rect x="250" y="80" width="820" height="560" rx="12" class="region"/>
  <text x="265" y="100" fill="#fbbf24" font-size="10" font-weight="600">应用服务区 / Docker Compose</text>
  {box(40,135,160,60,"个人用户","Web / Mobile H5","rgba(30,41,59,0.5)","#94a3b8")}
  {box(40,245,160,60,"医生/健康顾问","报告与风险查看","rgba(30,41,59,0.5)","#94a3b8")}
  {box(40,355,160,60,"平台管理员","后台管理","rgba(30,41,59,0.5)","#94a3b8")}
  {box(285,135,180,70,"Nginx","静态资源 / 反向代理","rgba(120,53,15,0.3)","#fbbf24")}
  {box(285,270,180,70,"Vue 3 管理后台","看板 / 表单 / 报表","rgba(8,51,68,0.4)","#22d3ee")}
  {box(515,120,185,65,"API Gateway","JWT / 限流 / 路由","rgba(6,78,59,0.4)","#34d399")}
  {box(515,215,185,65,"业务服务","用户 档案 指标 报告","rgba(6,78,59,0.4)","#34d399")}
  {box(515,310,185,65,"分析预警服务","趋势计算 / 规则引擎","rgba(6,78,59,0.4)","#34d399")}
  {box(515,405,185,65,"通知服务","站内提醒 / 邮件预留","rgba(6,78,59,0.4)","#34d399")}
  {cylinder(770,120,"PostgreSQL","业务数据 / 审计")}
  {cylinder(770,245,"Redis","缓存 / 会话 / 队列")}
  {cylinder(770,370,"MinIO","报告 / 体检附件")}
  {box(955,170,150,60,"备份与监控","日志 / 健康检查","rgba(120,53,15,0.3)","#fbbf24")}
  <path class="line2" d="M200,165 L285,165"/>
  <path class="line2" d="M200,275 L285,305"/>
  <path class="line2" d="M200,385 L285,305"/>
  <path class="line2" d="M465,168 L515,150"/>
  <path class="line2" d="M465,305 L515,248"/>
  <path class="line" d="M700,150 L770,150"/>
  <path class="line" d="M700,248 L770,150"/>
  <path class="line" d="M700,343 L770,275"/>
  <path class="line" d="M700,438 L770,400"/>
  <path class="line" d="M915,150 L955,200"/>
  <path class="line" d="M915,275 L955,200"/>
  <path class="line" d="M915,400 L955,200"/>
  <rect x="510" y="535" width="430" height="28" rx="8" fill="rgba(251,146,60,0.3)" stroke="#fb923c"/>
  <text x="725" y="553" class="label">异步任务通道：报告生成、导入解析、通知重试</text>
  <path class="line" d="M608,470 L608,535"/>
  <path class="line" d="M795,535 L795,440"/>
'''

flowchart = '''  <text x="30" y="36" class="title">个人健康生活网络数据监测与分析系统｜核心业务流程图</text>
  <text x="30" y="56" class="sub">从数据录入到预警处理、报告生成与计划跟进的闭环</text>
  <rect x="90" y="90" width="210" height="42" rx="21" fill="rgba(59,130,246,0.3)" stroke="#60a5fa" stroke-width="1.5"/>
  <text x="195" y="116" class="label">用户登录并选择档案</text>
  <rect x="90" y="170" width="210" height="54" rx="6" fill="#0f172a"/><rect x="90" y="170" width="210" height="54" rx="6" fill="rgba(8,51,68,0.4)" stroke="#22d3ee" stroke-width="1.5"/><text x="195" y="202" class="label">录入 / 导入健康数据</text>
  <g transform="translate(195,290)"><polygon points="0,-42 62,0 0,42 -62,0" fill="#0f172a"/><polygon points="0,-42 62,0 0,42 -62,0" fill="rgba(120,53,15,0.3)" stroke="#fbbf24" stroke-width="1.5"/><text y="-4" class="label">数据有效?</text><text y="14" class="small">范围/单位/时间</text></g>
  <rect x="440" y="263" width="215" height="54" rx="6" fill="#0f172a"/><rect x="440" y="263" width="215" height="54" rx="6" fill="rgba(136,19,55,0.4)" stroke="#fb7185" stroke-width="1.5"/><text x="547.5" y="296" class="label">提示修正并保留草稿</text>
  <rect x="90" y="380" width="210" height="54" rx="6" fill="#0f172a"/><rect x="90" y="380" width="210" height="54" rx="6" fill="rgba(6,78,59,0.4)" stroke="#34d399" stroke-width="1.5"/><text x="195" y="413" class="label">写入健康记录</text>
  <g transform="translate(195,510)"><polygon points="0,-42 62,0 0,42 -62,0" fill="#0f172a"/><polygon points="0,-42 62,0 0,42 -62,0" fill="rgba(120,53,15,0.3)" stroke="#fbbf24" stroke-width="1.5"/><text y="-4" class="label">触发风险?</text><text y="14" class="small">阈值/组合规则</text></g>
  <rect x="440" y="483" width="215" height="54" rx="6" fill="#0f172a"/><rect x="440" y="483" width="215" height="54" rx="6" fill="rgba(136,19,55,0.4)" stroke="#fb7185" stroke-width="1.5"/><text x="547.5" y="505" class="label">生成风险事件</text><text x="547.5" y="523" class="small">提醒用户/医生</text>
  <rect x="90" y="610" width="210" height="54" rx="6" fill="#0f172a"/><rect x="90" y="610" width="210" height="54" rx="6" fill="rgba(6,78,59,0.4)" stroke="#34d399" stroke-width="1.5"/><text x="195" y="642" class="label">更新趋势与健康总览</text>
  <rect x="440" y="610" width="215" height="54" rx="6" fill="#0f172a"/><rect x="440" y="610" width="215" height="54" rx="6" fill="rgba(6,78,59,0.4)" stroke="#34d399" stroke-width="1.5"/><text x="547.5" y="642" class="label">生成报告和调整计划</text>
  <rect x="775" y="610" width="185" height="42" rx="21" fill="rgba(59,130,246,0.3)" stroke="#60a5fa" stroke-width="1.5"/><text x="867.5" y="636" class="label">完成闭环跟进</text>
  <path class="line2" d="M195,132 L195,170"/><path class="line2" d="M195,224 L195,248"/><path class="line2" d="M195,332 L195,380"/><text x="210" y="360" fill="#34d399" font-size="9">是</text>
  <path class="line" d="M257,290 L440,290"/><text x="330" y="282" fill="#fb7185" font-size="9">否</text><path class="line" d="M547,263 C547,195 320,195 300,195"/>
  <path class="line2" d="M195,434 L195,468"/><path class="line2" d="M195,552 L195,610"/><text x="210" y="585" fill="#94a3b8" font-size="9">否</text>
  <path class="line" d="M257,510 L440,510"/><text x="330" y="502" fill="#fb7185" font-size="9">是</text>
  <path class="line" d="M547,537 L547,610"/><path class="line2" d="M300,637 L440,637"/><path class="line2" d="M655,637 L775,631"/>
'''

dataflow = f'''  <text x="30" y="36" class="title">个人健康生活网络数据监测与分析系统｜数据流图</text>
  <text x="30" y="56" class="sub">数据来源、处理链路、存储、分析和对外输出</text>
  {box(50,130,170,58,"手动录入","血压/血糖/体重等","rgba(30,41,59,0.5)","#94a3b8")}
  {box(50,245,170,58,"CSV / 体检附件","批量数据与文件","rgba(30,41,59,0.5)","#94a3b8")}
  {box(50,360,170,58,"设备接口预留","可穿戴数据源","rgba(30,41,59,0.5)","#94a3b8")}
  {box(310,180,190,62,"数据校验层","单位/范围/重复/权限","rgba(8,51,68,0.4)","#22d3ee")}
  {box(310,330,190,62,"导入解析任务","异步解析与失败记录","rgba(251,146,60,0.3)","#fb923c")}
  {cylinder(610,135,"PostgreSQL","结构化健康数据")}
  {cylinder(610,285,"MinIO","报告/附件对象")}
  {cylinder(610,435,"Redis","缓存/任务状态")}
  {box(850,140,190,62,"趋势分析","周期统计/异常点","rgba(6,78,59,0.4)","#34d399")}
  {box(850,285,190,62,"风险评估","阈值与组合规则","rgba(136,19,55,0.4)","#fb7185")}
  {box(850,430,190,62,"报告生成","周报/月报/PDF","rgba(6,78,59,0.4)","#34d399")}
  {box(1060,210,145,58,"看板输出","用户/后台","rgba(59,130,246,0.3)","#60a5fa")}
  {box(1060,380,145,58,"通知输出","提醒/处理","rgba(59,130,246,0.3)","#60a5fa")}
  <path class="line2" d="M220,159 L310,205"/><path class="line2" d="M220,274 L310,361"/><path class="line" d="M220,389 L310,361"/>
  <path class="line2" d="M500,211 L610,165"/><path class="line" d="M500,361 L610,315"/><path class="line" d="M500,361 L610,465"/>
  <path class="line2" d="M755,165 L850,171"/><path class="line2" d="M755,165 L850,316"/><path class="line" d="M755,315 L850,461"/><path class="line" d="M755,465 L850,461"/>
  <path class="line2" d="M1040,171 L1060,239"/><path class="line" d="M1040,316 L1060,409"/><path class="line2" d="M1040,461 L1060,409"/>
  <rect x="295" y="555" width="640" height="44" rx="8" fill="rgba(120,53,15,0.3)" stroke="#fbbf24"/>
  <text x="615" y="579" class="label">隐私与审计：授权范围、访问日志、导出记录、数据脱敏策略贯穿全链路</text>
'''

write_svg(DIAGRAM_DIR / "system-architecture.svg", architecture)
write_svg(DIAGRAM_DIR / "business-flow.svg", flowchart, "0 0 1000 720")
write_svg(DIAGRAM_DIR / "data-flow.svg", dataflow, "0 0 1240 660")


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color="FFFFFF")
        set_cell_shading(hdr[i], "1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    doc.add_paragraph()
    return table


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Arial"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if level == 1:
            r.font.color.rgb = RGBColor(46, 116, 181)
        else:
            r.font.color.rgb = RGBColor(31, 77, 120)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        rest = text[len(bold_prefix):]
        p.add_run(rest)
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.name = "Arial"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run(project["name"])
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(22)
r.bold = True
r.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("需求规格说明书 · 概要设计 · 详细设计 · 测试与验收文档")
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(67, 67, 67)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("版本：V1.0  |  适用阶段：立项、开发、联调、验收  |  技术栈：" + project["stack"])
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(9)

add_heading(doc, "1. 项目概述", 1)
add_para(doc, project["summary"])
add_para(doc, "系统以健康档案和连续指标数据为核心，形成“采集、清洗、分析、预警、报告、计划跟进”的产品闭环。平台同时提供个人端和运营后台，支持从 MVP 到可商业化运营的渐进建设。")

add_heading(doc, "2. 建设目标", 1)
for item in [
    "建立统一健康档案和多指标数据模型，支持长期、可追溯的数据沉淀。",
    "提供趋势分析、异常预警和周期报告，帮助用户理解健康变化。",
    "建设后台规则、字典、用户、文件和审计能力，保障可运营、可管理。",
    "用容器化部署和可观测性设计降低交付和运维成本。",
]:
    add_para(doc, item)

add_heading(doc, "3. 用户角色", 1)
add_table(doc, ["角色", "主要诉求", "关键权限"], [
    ("个人用户", "记录健康数据、查看趋势、接收提醒、导出报告", "本人档案与数据管理"),
    ("家庭管理员", "查看授权家庭成员健康状态", "被授权成员数据查看与提醒"),
    ("医生/健康顾问", "查看用户授权报告与风险摘要", "授权范围内的只读访问与建议备注"),
    ("平台运营", "维护用户、规则、字典、通知和文件", "后台业务管理"),
    ("系统管理员", "配置角色权限、系统参数、审计与备份", "最高管理权限"),
])

add_heading(doc, "4. 功能模块清单", 1)
add_table(doc, ["模块", "功能说明"], modules)

add_heading(doc, "5. 页面清单", 1)
add_table(doc, ["端", "页面", "页面职责"], pages)

add_heading(doc, "6. 数据库表设计", 1)
add_table(doc, ["表名", "中文名", "核心字段"], tables)
add_para(doc, "数据库建议采用 PostgreSQL。健康指标记录按 user_id、metric_code、measured_at 建联合索引；审计、通知、风险事件按 created_at 建时间索引；文件表只保存对象元数据，实际文件存储在 MinIO。")

add_heading(doc, "7. API 接口规划", 1)
add_table(doc, ["方法", "路径", "用途"], apis)
add_para(doc, "接口统一使用 /api/v1 前缀，认证采用 JWT + refresh token。后台接口需 RBAC 权限校验，用户端数据接口必须校验本人或授权范围。")

add_heading(doc, "8. 前端开发计划", 1)
add_para(doc, "前端建议采用 Vue 3 + TypeScript + Vite + Pinia + Vue Router + ECharts + Element Plus。")
for item in [
    "第 1 阶段：搭建路由、权限守卫、布局框架、设计 token、接口请求封装。",
    "第 2 阶段：完成健康总览、指标录入、趋势分析、报告列表、家庭成员页面。",
    "第 3 阶段：完成后台用户、指标字典、预警规则、文件管理、审计日志页面。",
    "第 4 阶段：完成移动端响应式、空状态、错误状态、骨架屏、表单校验和可访问性检查。",
]:
    add_para(doc, item)
add_para(doc, "界面方向：采用 frontend-design 的 Swiss 信息密度风格，白底、清晰网格、单一蓝色强调；后台按 ui-ux-pro-max 要求保证表单标签可见、触控目标不低于 44px、移动端无横向滚动、数据图表不只依赖颜色表达。")

add_heading(doc, "9. 后端开发计划", 1)
add_para(doc, "后端建议采用 FastAPI + SQLAlchemy + Alembic + Pydantic + Celery/RQ + Redis。FastAPI 适合本项目的数据分析、规则计算、异步任务和 Python 生态扩展。")
for item in [
    "领域建模：用户、档案、指标、记录、规则、风险、报告、文件、通知、审计。",
    "基础能力：认证授权、RBAC、数据库迁移、统一异常、日志、OpenAPI 文档。",
    "业务能力：健康数据 CRUD、导入解析、趋势聚合、规则触发、报告生成。",
    "运维能力：健康检查、限流、任务重试、对象存储签名上传、审计留痕。",
]:
    add_para(doc, item)

add_heading(doc, "10. 测试计划", 1)
add_table(doc, ["测试类型", "范围", "通过标准"], [
    ("单元测试", "规则引擎、指标校验、权限判断、报告摘要计算", "核心模块覆盖率不低于 80%"),
    ("接口测试", "认证、档案、指标、风险、报告、后台接口", "主要接口 100% 有正反用例"),
    ("前端测试", "表单校验、路由权限、图表筛选、空/错/加载状态", "关键页面通过组件和端到端测试"),
    ("性能测试", "健康记录查询、趋势聚合、批量导入、报告生成", "常用接口 P95 小于 500ms，导入任务可异步完成"),
    ("安全测试", "越权访问、导出权限、文件下载、敏感日志", "无高危漏洞，权限边界可验证"),
    ("验收测试", "按用户角色完成端到端业务闭环", "所有 P0/P1 验收用例通过"),
])

add_heading(doc, "11. 部署计划", 1)
for item in [
    "使用 Docker Compose 编排 nginx、frontend、api、worker、postgres、redis、minio。",
    "Nginx 负责 HTTPS、静态资源、接口反向代理和上传大小限制。",
    "数据库每日备份，MinIO 开启对象版本管理或定期快照，敏感配置通过环境变量注入。",
    "上线流程为测试环境联调、预发布验收、生产发布、发布后 24 小时监控。",
]:
    add_para(doc, item)

add_heading(doc, "12. 验收标准", 1)
for item in [
    "个人用户可完成注册登录、档案维护、指标录入、趋势查看、报告导出和风险提醒查看。",
    "后台可维护用户、角色、指标字典、预警规则、文件和审计日志。",
    "规则触发、通知记录、报告生成和权限授权流程可被测试用例验证。",
    "部署文档完整，容器化一键启动，关键服务有健康检查。",
    "性能、安全、兼容性和数据备份达到约定指标。",
]:
    add_para(doc, item)

add_heading(doc, "13. 开发里程碑", 1)
add_table(doc, ["里程碑", "周期", "交付物"], milestones)

add_heading(doc, "14. 人员分工建议", 1)
add_table(doc, ["角色", "人数", "职责"], [
    ("项目经理/产品负责人", "1", "需求澄清、排期、风险、验收组织"),
    ("UI/UX 设计师", "1", "原型、设计系统、响应式与后台体验"),
    ("前端工程师", "2", "用户端、后台、图表、前端测试"),
    ("后端工程师", "2", "API、权限、规则、报告、任务、存储"),
    ("测试工程师", "1", "测试计划、用例、自动化、验收回归"),
    ("运维/DevOps", "0.5", "Docker、Nginx、CI/CD、备份和监控"),
])

add_heading(doc, "15. 风险与应对措施", 1)
add_table(doc, ["风险", "影响", "应对措施"], [
    ("健康数据隐私合规不足", "影响上线与用户信任", "最小权限、审计、脱敏、授权记录、导出留痕"),
    ("指标范围和预警规则不准确", "误报或漏报", "规则字典可配置，先用提示型预警，邀请专业人员评审"),
    ("设备接口不稳定", "数据接入延期", "MVP 先支持手动和 CSV，设备接口作为扩展适配层"),
    ("趋势分析性能不足", "看板响应慢", "聚合缓存、分页、时间索引、异步任务"),
    ("报告导出格式复杂", "交付体验不稳定", "先固定模板，异步生成，失败可重试并记录原因"),
    ("需求范围膨胀", "延期", "按 P0/P1/P2 管理范围，里程碑验收冻结变更"),
])

add_heading(doc, "附录：图示文件", 1)
for name in ["system-architecture.svg", "business-flow.svg", "data-flow.svg"]:
    add_para(doc, f"{name} 已输出至 deliverables/diagrams/personal-health-platform。")

doc_path = BASE / "个人健康生活网络数据监测与分析系统_项目交付文档.docx"
doc.save(doc_path)
print(doc_path)
